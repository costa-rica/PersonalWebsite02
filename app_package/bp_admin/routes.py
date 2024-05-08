
from flask import Blueprint
from flask import render_template, url_for, redirect, flash, request, \
    abort, session, Response, current_app, send_from_directory, make_response, \
    g, after_this_request
import bcrypt
from flask_login import login_required, login_user, logout_user, current_user
import os
import json
from pw_models import Base, engine, DatabaseSession, text, Users
# from app_package.bp_admin.utils import get_user_loc_day_tuple
import pandas as pd
import shutil
from datetime import datetime, timedelta
import openpyxl
import zipfile
# from ws_utilities import create_df_crosswalk, update_and_append_user_location_day, \
#     update_and_append_via_df_crosswalk_users, update_and_append_via_df_crosswalk_locations
# from ws_utilities import create_df_from_db_table_name, request_visual_crossing_for_one_day, \
#     request_visual_crossing_for_last_30days, add_weather_history
from app_package._common.utilities import custom_logger, wrap_up_session


logger_bp_admin = custom_logger('bp_admin.log')
# salt = bcrypt.gensalt()
bp_admin = Blueprint('bp_admin', __name__)


# @bp_admin.before_request
# def before_request():
#     logger_bp_admin.info(f"- in before_request route --")
#     if not current_user.admin_users_permission:
#         return redirect(url_for('bp_users.user_home'))
    
#     g.db_session = DatabaseSession()

#     if request.referrer:
#         logger_bp_admin.info(f"- request.referrer: {request.referrer} ")
    
#     logger_bp_admin.info(f"- db_session ID: {id(g.db_session)} ")
    
#     if request.endpoint:
#         logger_bp_admin.info(f"- request.endpoint: {request.endpoint} ")

    
#     # TEMPORARILY_DOWN: redirects to under construction page #
#     if os.environ.get('TEMPORARILY_DOWN') == '1':
#         if request.url != request.url_root + url_for('bp_main.temporarily_down')[1:]:
#             logger_bp_admin.info(f'- request.referrer: {request.referrer}')
#             logger_bp_admin.info(f'- request.url: {request.url}')
#             return redirect(url_for('bp_main.temporarily_down'))


# @bp_admin.after_request
# def after_request(response):
#     logger_bp_admin.info(f"---- after_request --- ")
#     if hasattr(g, 'db_session'):
#         wrap_up_session(logger_bp_admin, g.db_session)
#     return response


@bp_admin.before_request
def before_request():
    logger_bp_admin.info("-- def before_request() --")
    # Assign a new session to a global `g` object, accessible during the whole request
    g.db_session = DatabaseSession()
    
    # Use getattr to safely access g.referrer, defaulting to None if it's not set
    if getattr(g, 'referrer', None) is None:
        if request.referrer:
            g.referrer = request.referrer
        else:
            g.referrer = "No referrer"
    
    logger_bp_admin.info("-- def before_request() END --")


@bp_admin.route('/admin_page', methods = ['GET', 'POST'])
@login_required
def admin_page():
    db_session = g.db_session




    # if request.method == 'POST':
    #     formDict = request.form.to_dict()
    #     print('formDict:::', formDict)
    #     # if formDict.get('add_link'):
    #     if 'add_link' in formDict.keys():
    #         print("**** add ing link ****")
    #         if formDict.get('input_test_flight_link')[:8]=="https://":
    #             with open(os.path.join(current_app.config.get('WEBSITE_FILES'),'TestFlightUrl.txt'), 'w') as file:
    #                 file.write(formDict.get('input_test_flight_link'))
    #             flash(f'Updated TestFlight link', 'success')
    #         else:
    #             flash(f'TestFlight link must be a valid https url', 'warning')    
    #         return redirect(url_for('bp_admin.admin_page'))

    #     # elif formDict.get('delete_link'):
    #     elif "delete_link" in formDict.keys():
    #         print("**** deleting link ****")
    #         os.remove(os.path.join(current_app.config.get('WEBSITE_FILES'),'TestFlightUrl.txt'))
    #         test_flight_link=""
    #         # flash(f'Removed TestFlight link from {request.url}', 'warning')
    #         flash(f'Removed TestFlight link from {request.host}', 'warning')


    return render_template('admin/admin_home.html', str=str)


@bp_admin.route('/admin_db_download', methods = ['GET', 'POST'])
@login_required
def admin_db_download():
    logger_bp_admin.info('- in admin_db_download -')
    # logger_bp_admin.info(f"current_user.admin_users_permission: {current_user.admin_users_permission}")
    db_session = g.db_session

    # if not current_user.admin_users_permission:
    #     return redirect(url_for('bp_main.home'))

    metadata = Base.metadata
    db_table_list = [table for table in metadata.tables.keys()]

    # csv_dir_path = os.path.join(current_app.config.get('DATABASE_HELPER_FILES'), 'db_backup')
    db_tables_dir_path = current_app.config.get('DIR_DB_DOWNLOAD')

    # delete the existing .zip file
    if os.path.exists(os.path.join(current_app.config.get('DATABASE_ROOT'),"db_download.zip")):
        os.remove(os.path.join(current_app.config.get('DATABASE_ROOT'),"db_download.zip"))

    if request.method == "POST":
        formDict = request.form.to_dict()
        if os.path.exists(os.path.join(current_app.config.get('DATABASE_ROOT'),"db_download")):
            for filename in os.listdir(os.path.join(current_app.config.get('DATABASE_ROOT'),"db_download")):
                os.remove(os.path.join(current_app.config.get('DATABASE_ROOT'),"db_download",filename))


        db_table_list = []
        for key, value in formDict.items():
            if value == "db_table":
                db_table_list.append(key)
      
        db_tables_dict = {}
        for table_name in db_table_list:
            # base_query = sess.query(metadata.tables[table_name])
            base_query = db_session.query(metadata.tables[table_name])
            df = pd.read_sql(text(str(base_query)), engine.connect())

            # fix table names
            cols = list(df.columns)
            for col in cols:
                if col[:len(table_name)] == table_name:
                    df = df.rename(columns=({col: col[len(table_name)+1:]}))


            db_tables_dict[table_name] = df
            if formDict.get("download_files") == "csv":
                # db_tables_dict[table_name].to_csv(os.path.join(csv_dir_path, f"{table_name}.csv"), index=False)
                db_tables_dict[table_name].to_csv(os.path.join(db_tables_dir_path, f"{table_name}.csv"), index=False)
            elif formDict.get("download_files") == "pickle":
                # db_tables_dict[table_name].to_pickle(os.path.join(csv_dir_path, f"{table_name}.pkl"))
                db_tables_dict[table_name].to_pickle(os.path.join(db_tables_dir_path, f"{table_name}.pkl"))
        
        # shutil.make_archive(csv_dir_path, 'zip', csv_dir_path)
        shutil.make_archive(db_tables_dir_path, 'zip', db_tables_dir_path)
        logger_bp_admin.info(f"--- about to download db_download.zip ----")
        # return redirect(url_for('bp_admin.download_db_tables_zip'))


        # Ensure the file is deleted after the request
        @after_this_request
        def remove_file(response):
            try:
                os.remove(os.path.join(current_app.config['DATABASE_ROOT'],'db_download.zip'))
            except Exception as error:
                app.logger.error("Error removing or closing downloaded file handle", error)
            return response

        
        return send_from_directory(os.path.join(current_app.config['DATABASE_ROOT']),'db_download.zip', as_attachment=True)
    
    return render_template('admin/admin_db_download_page.html', db_table_list=db_table_list )


@bp_admin.route('/admin_db_upload_single_file', methods = ['GET', 'POST'])
@login_required
def admin_db_upload_single_file():
    logger_bp_admin.info('- in admin_db_upload_single_file -')

    metadata = Base.metadata
    db_table_list = [table for table in metadata.tables.keys()]
    logger_bp_admin.info(f'- Table List -')
    for table in db_table_list:
        logger_bp_admin.info(f'- {table} -')

    list_files_in_db_upload = os.listdir(current_app.config.get('DIR_DB_UPLOAD'))
    list_files_in_db_upload_csv_pkl_zip = []
    for filename_string in list_files_in_db_upload:
        filename, file_extension = os.path.splitext(filename_string)
        if file_extension in ['.csv','.pkl']:
            list_files_in_db_upload_csv_pkl_zip.append(filename_string)

    if request.method == "POST":
        formDict = request.form.to_dict()
        print(f"- admin_db_upload_single_file POST -")
        print("formDict: ", formDict)


        if formDict.get('what_kind_of_post') == "upload_from_here":
            requestFiles = request.files

            
            # csv_file_for_table = request.files.get('csv_table_upload')
            file_for_table_upload = request.files.get('file_for_table_upload')
            file_for_table_upload_filename = file_for_table_upload.filename

            logger_bp_admin.info(f"-- Get CSV file name --")
            logger_bp_admin.info(f"--  {file_for_table_upload_filename} --")

            ## save to databases/personalWEb02/db_upload directory
            path_to_uploaded_table_file = os.path.join(current_app.config.get('DIR_DB_UPLOAD'),file_for_table_upload_filename)
            file_for_table_upload.save(path_to_uploaded_table_file)

            logger_bp_admin.info(f"-- table to go to: { formDict.get('existing_db_table_to_update')}")

            return redirect(url_for('bp_admin.upload_table', table_name = formDict.get('existing_db_table_to_update'),
                path_to_uploaded_table_file=path_to_uploaded_table_file))
                # path_to_uploaded_csv=path_to_uploaded_csv))
    #     elif formDict.get('what_kind_of_post') == "uploaded_already":
    #         already_uploaded_filename = formDict.get('selectedFile')
    #         path_to_uploaded_table_file = os.path.join(current_app.config.get('DIR_DB_UPLOAD'),already_uploaded_filename)
            
    #         return redirect(url_for('bp_admin.upload_table', table_name = formDict.get('existing_db_table_to_update'),
    #             path_to_uploaded_table_file=path_to_uploaded_table_file))

    return render_template('admin/admin_db_upload_single_file_page.html', db_table_list=db_table_list,
        len=len, list_files_in_db_upload_csv_pkl_zip=list_files_in_db_upload_csv_pkl_zip)


@bp_admin.route('/upload_table/<table_name>', methods = ['GET', 'POST'])
@login_required
def upload_table(table_name):
    logger_bp_admin.info('- in upload_table -')

    db_session = g.db_session


    path_to_uploaded_table_file = request.args.get('path_to_uploaded_table_file')


    # Get Table Column names from the database corresponding to the running webiste/app
    metadata = Base.metadata
    existing_table_column_names = metadata.tables[table_name].columns.keys()

    if path_to_uploaded_table_file[-4:]=='.pkl':
        df = pd.read_pickle(path_to_uploaded_table_file)
    else:
        # Get column names from the uploaded csv
        df = pd.read_csv(path_to_uploaded_table_file)

    if 'time_stamp_utc' in df.columns and path_to_uploaded_table_file[-4:]!='.pkl':
        try:
            df['time_stamp_utc'] = pd.to_datetime(df['time_stamp_utc'], format='%d/%m/%Y %H:%M')
        except:
            df = pd.read_csv(path_to_uploaded_table_file, parse_dates=['time_stamp_utc'])

    replacement_data_col_names = list(df.columns)
    
    # Match column names between the two tables
    match_cols_dict = {}
    for existing_db_column in existing_table_column_names:
        try:
            index = replacement_data_col_names.index(existing_db_column)
            match_cols_dict[existing_db_column] = replacement_data_col_names[index]
        except ValueError:
            match_cols_dict[existing_db_column] = None

    # if request.method == "POST":
    #     formDict = request.form.to_dict()


    #     # NOTE: upload data to existing database
    #     ### formDict (key) is existing databaes column name
    #     # existing_names_list = [existing for existing, update in formDict.items() if update != 'true' ]
        
    #     # check for default values and remove from formDict
    #     set_default_value_dict = {}
    #     for key, value in formDict.items():
    #         if key[:len("default_checkbox_")] == "default_checkbox_":
    #             set_default_value_dict[value] = formDict.get(value)
        

    #     # Delete elements from dictionary
    #     for key, value in set_default_value_dict.items():
    #         del formDict[key]
    #         checkbox_key = "default_checkbox_" + key
    #         del formDict[checkbox_key]

    #     existing_names_list = []
    #     for key, value in formDict.items():
    #         if value != 'true':
    #             existing_names_list.append(key)
            
    #     df_update = pd.DataFrame(columns=existing_names_list)

    #     # value is the new data (aka the uploaded csv file column)
    #     for exisiting, replacement in formDict.items():
    #         if not replacement in ['true','']:
    #             # print(replacement)
    #             df_update[exisiting]=df[replacement].values

    #     # Add in columns with default values
    #     for column_name, default_value in set_default_value_dict.items():
    #         if column_name == 'time_stamp_utc': 
    #             df_update[column_name] = datetime.utcnow()
    #         else:
    #             df_update[column_name] = default_value

        
    #     # remove existing users from upload
    #     # NOTE: There needs to be a user to upload data
    #     if table_name == 'users':
    #         print("--- Found users table ---")
    #         existing_users = db_session.query(Users).all()
    #         list_of_emails_in_db = [i.email for i in existing_users]
    #         for email in list_of_emails_in_db:
    #             # df_update.drop(df_update[df_update.email== email].index, inplace = True)
    #             print(f"-- removing {email} from upload dataset --")
    #             df_update.drop(df_update[df_update.email == email].index, inplace=True)
    #             df_update.reset_index(drop=True, inplace=True)

    #         # Assuming 'password' column exists and you want to encode all non-null passwords
    #         if 'password' in df_update.columns:
    #             df_update['password'] = df_update['password'].apply(lambda x: x.encode() if pd.notnull(x) else x)

    #     df_update.to_sql(table_name, con=engine, if_exists='append', index=False)
    #     # wrap_up_session(logger_bp_admin)
    #     flash(f"{table_name} update: successful!", "success")

    #     # return redirect(request.url)
    #     return redirect(url_for('bp_admin.admin_db_upload_single_file'))
    
    return render_template('admin/upload_table.html', table_name=table_name, 
        match_cols_dict = match_cols_dict,
        existing_table_column_names=existing_table_column_names,
        replacement_data_col_names = replacement_data_col_names)


# @bp_admin.route('/admin_db_upload_zip', methods = ['GET', 'POST'])
# @login_required
# def admin_db_upload_zip():
#     logger_bp_admin.info('- in admin_db_upload_zip -')

#     if not current_user.admin_users_permission:
#         return redirect(url_for('bp_main.home'))

#     metadata = Base.metadata
#     db_table_list = [table for table in metadata.tables.keys()]
    
#     list_files_in_db_upload = os.listdir(current_app.config.get('DB_UPLOAD'))
#     list_files_in_db_upload_csv_pkl_zip = []
#     for filename_string in list_files_in_db_upload:
#         filename, file_extension = os.path.splitext(filename_string)

#         if file_extension in ['.zip']:
#             list_files_in_db_upload_csv_pkl_zip.append(filename_string)

#     if request.method == "POST":
#         formDict = request.form.to_dict()
#         logger_bp_admin.info(f"- admin_db_upload_zip POST -")
#         logger_bp_admin.info(f"formDict: {formDict}")
#         logger_bp_admin.info(f"request.files: {request.files}")
        
#         file_for_table_upload = request.files.get('zip_filename_uploaded')
        


#         if file_for_table_upload.filename != '':

#             zip_filename = file_for_table_upload.filename

#             logger_bp_admin.info(f"-- Get .zip file name --")
#             logger_bp_admin.info(f"-- {zip_filename} --")

#             ## save to databases/WhatSticks/database_helpers/DB_UPLOAD directory
#             path_to_uploaded_table_file = os.path.join(current_app.config.get('DB_UPLOAD'),zip_filename)
#             file_for_table_upload.save(path_to_uploaded_table_file)
#         else:
#             zip_filename = formDict.get('zip_filename_existing')

#         if zip_filename == None:
#             flash(f"Select a .zip file", "warning")
#             return redirect(request.referrer)
#         count_of_new_users = 0
#         count_of_new_locations = 0
#         count_of_new_user_location_day = 0
#         count_of_new_weather_hist = 0
#         count_of_new_workouts = 0
#         count_of_new_qty_cat = 0

#         # Step 1: Make Crosswalks 
#         df_crosswalk_users = create_df_crosswalk('users', zip_filename)
#         df_crosswalk_locations = create_df_crosswalk('locations', zip_filename)

#         logger_bp_admin.info(f"-- count of df_crosswalk_users {len(df_crosswalk_users)} --")
#         logger_bp_admin.info(f"-- count of df_crosswalk_locations {len(df_crosswalk_locations)} --")

#         if 'new_row' in df_crosswalk_users.columns:
#             count_of_new_users = len(df_crosswalk_users[df_crosswalk_users.new_row=='yes'])
#             df_crosswalk_users.drop(columns=['new_row'], inplace=True)
#         if 'new_row' in df_crosswalk_locations.columns:
#             count_of_new_locations = len(df_crosswalk_locations[df_crosswalk_locations.new_row=='yes'])
#             df_crosswalk_locations.drop(columns=['new_row'], inplace=True)

#         if len(df_crosswalk_locations) > 0 and len(df_crosswalk_users) > 0:
#             # Step 2: Add UserLocationDay data with new user_ids and location_ids, if any new rows to add
#             count_of_new_user_location_day = update_and_append_user_location_day(
#                 zip_filename,df_crosswalk_users,df_crosswalk_locations)
        
#         if len(df_crosswalk_locations) > 0:
#             # Step 3: Add WeatherHistory with new location_ids, if any new rows to add
#             count_of_new_weather_hist = update_and_append_via_df_crosswalk_locations(
#                 'weather_history', 'location_id', zip_filename,df_crosswalk_locations)

#         if len(df_crosswalk_users) > 0:
#             # Step 4: Add AppleHealthWorkouts with new user_ids, if any new rows to add
#             count_of_new_workouts = update_and_append_via_df_crosswalk_users(
#                 'apple_health_workout',zip_filename,df_crosswalk_users)
            
#             # Step 5: Add AppleHealthQuantityCategory with new user_ids, if any new rows to add
#             count_of_new_qty_cat = update_and_append_via_df_crosswalk_users(
#                 'apple_health_quantity_category',zip_filename,df_crosswalk_users)

#         # request.referrer - the url for the page that sent 
#         # in this case it's just this same page.

#         long_f_string = (
#             f"Successfully added: " +
#             f"\n Users....................... {count_of_new_users:,} " +
#             f"\n Locations................... {count_of_new_locations:,}" +
#             f"\n UserLocationDay............. {count_of_new_user_location_day:,}" +
#             f"\n WeatherHistory.............. {count_of_new_weather_hist:,}" +
#             f"\n Workouts.................... {count_of_new_workouts:,}" +
#             f"\n AppleHealthQuantityCategory. {count_of_new_qty_cat:,}"
#         )
#         # wrap_up_session(logger_bp_admin)
#         flash( long_f_string, "success")
#         return redirect(request.referrer)

#     return render_template('admin/admin_db_upload_zip_page.html', db_table_list=db_table_list,
#         len=len, list_files_in_db_upload_csv_pkl_zip=list_files_in_db_upload_csv_pkl_zip)



# # # @bp_admin.route("/download_db_tables_as_csv", methods=["GET","POST"])
# @bp_admin.route("/download_db_tables_zip", methods=["GET","POST"])
# @login_required
# def download_db_tables_zip():
#     # return send_from_directory(os.path.join(current_app.config['DATABASE_HELPER_FILES']),'db_backup.zip', as_attachment=True)
#     return send_from_directory(os.path.join(current_app.config['DATABASE_ROOT']),'db_download.zip', as_attachment=True)



@bp_admin.route("/delete_db_upload_file/<filename>", methods=["GET","POST"])
@login_required
def delete_db_upload_file(filename):
    logger_bp_admin.info(f"- accessed delete_db_upload_file -")
    logger_bp_admin.info(f"- deleting {filename} -")

    try:
        os.remove(os.path.join(current_app.config['DIR_DB_UPLOAD'],filename))
    except FileNotFoundError:
        logger_bp_admin.info(f"File {filename} was not found.")
    except PermissionError:
        logger_bp_admin.info(f"Permission denied: Unable to delete {filename}.")
    except Exception as e:
        logger_bp_admin.info(f"An error occurred: {e}")
    

    # Redirect back to the referrer page
    referrer = request.referrer
    if referrer:
        return redirect(referrer)
    else:
        # Fallback to a default route if the referrer is not available
        return redirect(url_for('bp_admin.admin_page'))



# @bp_admin.route("/delete_user/<email>", methods=["GET","POST"])
# @login_required
# def delete_user(email):
#     logger_bp_admin.info(f'-accessed: delete_user {email}')
#     # with open(os.path.join(current_app.config['DIR_DB_FILES_UTILITY'],'added_users.txt')) as json_file:
#     #     get_users_dict=json.load(json_file)
#     #     json_file.close()
    
#     # del get_users_dict[email]
    
#     # added_users_file=os.path.join(current_app.config['DIR_DB_FILES_UTILITY'], 'added_users.txt')
#     # with open(added_users_file, 'w') as json_file:
#     #     json.dump(get_users_dict, json_file)
        
#     # if len(sess.query(User).filter_by(email=email).all())>0:
#     #     sess.query(User).filter_by(email=email).delete()
#     #     sess_users.commit()
    
    
    
#     # flash(f'{email} has been deleted!', 'success')
#     flash(f'NOT deleteing emails with this link yet', 'success')
#     return redirect(url_for('bp_admin.admin_page'))
#     # return redirect(request.url)


# @bp_admin.route("/weather_history_admin", methods=["GET","POST"])
# @login_required
# def weather_history_admin():
#     logger_bp_admin.info(f'- in weather_history_admin ')
#     db_session = g.db_session
#     if not current_user.admin_users_permission:
#         return redirect(url_for('bp_main.home'))

#     # Get each weather history record
#     df_weather_history = create_df_from_db_table_name("weather_history")
#     df_locations = create_df_from_db_table_name("locations")
#     df_locations.rename(columns={'id': 'location_id'}, inplace=True)
#     df_weather_hist_with_locations = pd.merge(df_weather_history,
#                                           df_locations[["location_id","city","country"]],
#                                           on=["location_id"],how="left")
#     df_weather_hist_with_locations["date_time_obj"] = pd.to_datetime(df_weather_hist_with_locations.date_time, format="%Y-%m-%d")
#     df_sorted = df_weather_hist_with_locations.sort_values(by='date_time_obj', ascending=False)
#     dict_sorted = df_sorted.to_dict('records')

#     col_names = ["date", "location_id", "city","temp"]

#     # Get cities
#     # create GROUPBY dataframe with city location_id and count
#     df_city_location_counts = df_weather_hist_with_locations.groupby(['city', 'location_id']).size().reset_index(name='count')
#     dict_city_location_counts = df_city_location_counts.to_dict('records')

#     if request.method == "POST":
#         formDict = request.form.to_dict()
#         logger_bp_admin.info(f"formDict: {formDict}")
#         # add_location_weather_history = formDict.get("location_add_weather_history")
#         # add_location_weather_history_date = formDict.get("weather_hist_date")
#         add_weather_hist_location = formDict.get("add_weather_hist_location")# in the form <location_id: city>
#         if add_weather_hist_location:
#             print("**** getting id ******")
#             print("**** getting id ******")

#             add_weather_hist_location_id = add_weather_hist_location[:add_weather_hist_location.find(":")]
#             print(f"**** add_weather_hist_location_id: {add_weather_hist_location_id} ******")
#             print("**** getting id ******")
        
#         add_weather_hist_location_date_option = formDict.get("add_weather_hist_location_date_option")
#         add_weather_hist_location_date = formDict.get("add_weather_hist_location_date")
#         weather_record_id = formDict.get("delete_weather_record_id")
#         weather_hist_location_id = formDict.get("delete_weather_record_from_location_id")

#         if add_weather_hist_location:
#             location_to_add_weather = db_session.get(Locations, add_weather_hist_location_id)
            
#             if add_weather_hist_location_date_option == "yesterday":
#                 yesterday_date = datetime.utcnow()  - timedelta(days=1)
#                 vc_weather_dict = request_visual_crossing_for_one_day(
#                     location_to_add_weather,yesterday_date.strftime('%Y-%m-%d'))
#                 add_weather_history(db_session, add_weather_hist_location_id, vc_weather_dict)
#                 days_added = "ONE RECORD"

#             elif add_weather_hist_location_date_option == "last_thirty_days":
#                 vc_weather_dict = request_visual_crossing_for_last_30days(location_to_add_weather)
#                 days_added = "30 RECORDS"
#                 add_weather_history(db_session, add_weather_hist_location_id, vc_weather_dict)

#             elif add_weather_hist_location_date_option == "specific_date":
#                 vc_weather_dict = request_visual_crossing_for_one_day(location_to_add_weather,add_weather_hist_location_date)
#                 days_added = "ONE RECORD"
#                 add_weather_history(db_session, add_weather_hist_location_id, vc_weather_dict)
#             else:
#                 days_added = "NONE"

#             long_f_string = (
#                 f"Added {days_added} for \n" +
#                 f"id:{location_to_add_weather.id}, city: {location_to_add_weather.city} "
#             )
#             flash_banner_color = "success"

#         elif weather_record_id:
#             weather_record = db_session.get(WeatherHistory, weather_record_id)
#             weather_record_city = db_session.get(Locations, weather_record.location_id).city
#             db_session.query(WeatherHistory).filter_by(id=weather_record_id).delete()
#             long_f_string = (
#                 f"Deleted ONE RECORD for \n" +
#                 f"id:{weather_record_id}, city: {weather_record_city} "
#             )
#             flash_banner_color = "warning"

#         elif weather_hist_location_id:
#             # weather_record = db_session.get(WeatherHistory, weather_record_id)
#             # weather_record_city = db_session.get(Locations, weather_record.location_id).city
#             city_name = db_session.get(Locations,weather_hist_location_id).city
#             db_session.query(WeatherHistory).filter_by(location_id=weather_hist_location_id).delete()
#             long_f_string = (
#                 f"Deleted ALL RECORDS for \n" +
#                 f"location_id: {weather_hist_location_id}; city: {city_name}"
#             )
#             flash_banner_color = "warning"
        
#         else:
#             long_f_string = (
#                 f"No records affected"
#             )
#             flash_banner_color = "warning"

#         flash( long_f_string, flash_banner_color)
#         return redirect(request.referrer)

#     return render_template('admin/weather_history_admin.html', col_names = col_names,
#             dict_weather_hist_with_locations = dict_sorted, dict_city_location_counts = dict_city_location_counts)



################
###  OLD ###
################







# @bp_admin.route('/database_page', methods=["GET","POST"])
# @login_required
# def database_page():
#     tableNamesList=['investigations','tracking_inv','recalls','tracking_re','user']
#     # tableNamesList= db.engine.table_names()
#     legend='Database downloads'
#     if request.method == 'POST':
#         formDict = request.form.to_dict()
#         print('formDict::::', formDict)

#         if formDict.get('build_workbook')=="True":
            
#             #check if os.listdir(current_app.config['DIR_DB_FILES_DATABASE']), if no create:
#             if not os.path.exists(current_app.config['DIR_DB_FILES_DATABASE']):
#                 # print('There is not database folder found???')
#                 os.mkdir(current_app.config['DIR_DB_FILES_DATABASE'])
            
#             for file in os.listdir(current_app.config['DIR_DB_FILES_DATABASE']):
#                 os.remove(os.path.join(current_app.config['DIR_DB_FILES_DATABASE'], file))

            
#             timeStamp = datetime.now().strftime("%y%m%d_%H%M%S")
#             workbook_name=f"database_tables{timeStamp}.xlsx"
#             print('reportName:::', workbook_name)
#             excelObj=pd.ExcelWriter(os.path.join(current_app.config['DIR_DB_FILES_DATABASE'], workbook_name),
#                 date_format='yyyy/mm/dd', datetime_format='yyyy/mm/dd')
#             workbook=excelObj.book
            
#             dictKeyList=[i for i in list(formDict.keys()) if i in tableNamesList]
#             dfDictionary={h : pd.read_sql_table(h, db.engine) for h in dictKeyList}
#             for name, df in dfDictionary.items():
#                 if len(df)>900000:
#                     flash(f'Too many rows in {name} table', 'warning')
#                     return render_template('database.html',legend=legend, tableNamesList=tableNamesList)
#                 df.to_excel(excelObj,sheet_name=name, index=False)
#                 worksheet=excelObj.sheets[name]
#                 start_row=0
#                 formatExcelHeader(workbook,worksheet, df, start_row)
#                 print(name, ' table added to workbook')
#                 # if name=='dmrs':
#                     # dmrDateFormat = workbook.add_format({'num_format': 'yyyy-mm-dd'})
#                     # worksheet.set_column(1,1, 15, dmrDateFormat)
                
#             print('path of reports:::',os.path.join(current_app.config['DIR_DB_FILES_DATABASE'],str(workbook_name)))
#             excelObj.close()
#             print('excel object close')
#             # return send_from_directory(current_app.config['DIR_DB_FILES_DATABASE'],workbook_name, as_attachment=True)
#             return redirect(url_for('users.database_page'))

#         elif formDict.get('download_db_workbook'):
#             return redirect(url_for('users.download_db_workbook'))

#         elif formDict.get('uploadFileButton'):
#             # print('****uploadFileButton****')
#             logger_bp_admin.info("* upload excel file to ")
#             formDict = request.form.to_dict()
#             filesDict = request.files.to_dict()
#             # print('formDict:::',formDict)
#             # print('filesDict:::', filesDict)
            
            
#             if not os.path.exists(current_app.config['DIR_DB_FILES_TEMPORARY']):
#                 os.mkdir(current_app.config['DIR_DB_FILES_TEMPORARY'])
            
#             file_type=formDict.get('file_type')
#             uploadData=request.files['fileUpload']
#             uploadFileName=uploadData.filename
#             uploadData.save(os.path.join(current_app.config['DIR_DB_FILES_TEMPORARY'], uploadFileName))
#             if file_type=="excel":
#                 wb = openpyxl.load_workbook(uploadData)
#                 sheetNames=json.dumps(wb.sheetnames)
#                 tableNamesList=json.dumps(tableNamesList)

#                 return redirect(url_for('bp_admin.database_upload',legend=legend,tableNamesList=tableNamesList,
#                     sheetNames=sheetNames, uploadFileName=uploadFileName,file_type=file_type))
#             return redirect(url_for('bp_admin.database_upload',legend=legend,uploadFileName=uploadFileName,
#                 file_type=file_type))
#             # return redirect(url_for('users.database_page'))
#         elif formDict.get('btn_database_delete') and formDict.get('database_delete_verify') == 'delete':
#             logger_bp_admin.info("* Delete database")

            
#             sess.query(Investigations).delete()
#             sess.query(Tracking_inv).delete()
#             sess.query(Saved_queries_inv).delete()
#             sess.query(Recalls).delete()
#             sess.query(Tracking_re).delete()
#             sess.query(Saved_queries_re).delete()
#             sess_users.commit()

#             logger_bp_admin.info(f"- database (except users table) deleted by: {current_user.email}")
#             flash("Data tables (except for users table) successfully deleted", "warning")
#             return redirect(request.url)

#     return render_template('admin/database_page.html', legend=legend, tableNamesList=tableNamesList)



# @bp_admin.route("/download_db_workbook", methods=["GET","POST"])
# @login_required
# def download_db_workbook():
#     # workbook_name=request.args.get('workbook_name')
#     workbook_name = os.listdir(current_app.config['DIR_DB_FILES_DATABASE'])[0]
#     print('file:::', os.path.join(current_app.root_path, 'static','files_database'),workbook_name)
#     file_path = r'D:\OneDrive\Documents\professional\20210610kmDashboard2.0\fileShareApp\static\files_database\\'
    
#     return send_from_directory(os.path.join(current_app.config['DIR_DB_FILES_DATABASE']),workbook_name, as_attachment=True)


# @bp_admin.route('/database_upload', methods=["GET","POST"])
# @login_required
# def database_upload():
#     logger_bp_admin.info("- in database_upload route")
#     file_type=request.args.get('file_type')
#     if file_type=='excel':
#         tableNamesList=json.loads(request.args['tableNamesList'])
#         sheetNames=json.loads(request.args['sheetNames'])
#     uploadFileName=request.args.get('uploadFileName')
#     legend='Upload Data File to Database'
#     # uploadFlag=True
#     limit_upload_flag='checked'
    
#     if request.method == 'POST':
        
#         formDict = request.form.to_dict()
#         # print('formDict::::', formDict)
#         if formDict.get('appendExcel'):
            
#             uploaded_file=os.path.join(current_app.config['DIR_DB_FILES_TEMPORARY'], uploadFileName)
#             # print('uploaded_file::::',uploaded_file)
#             if file_type=='excel':
#                 logger_bp_admin.info("- in file_type=='excel'")
#                 sheet_upload_status = []
#                 for sheet in sheetNames:
#                     logger_bp_admin.info(f"- in sheet: {sheet}")
#                     sheetUpload=pd.read_excel(uploaded_file,engine='openpyxl',sheet_name=sheet)
#                     if sheet=='user':
#                         existing_emails=[i[0] for i in sess.query(Users.email).all()]
#                         sheetUpload=pd.read_excel(uploaded_file,engine='openpyxl',sheet_name='user')
#                         sheetUpload=sheetUpload[~sheetUpload['email'].isin(existing_emails)]

#                     elif formDict.get(sheet) in ['investigations','recalls']:
#                         sheetUpload['date_updated']=datetime.now()
#                         if formDict.get(sheet) =='recalls':
#                             sheetUpload=fix_recalls_wb_util(sheetUpload,uploadFileName)
#                         elif formDict.get(sheet) =='investigations':
#                             sheetUpload=fix_investigations_wb_util(sheetUpload)

#                     try:
#                         if sheet == 'user':
#                             sheetUpload.to_sql('users',con=engine, if_exists='append', index=False)
#                         elif sheet == 'recalls':
#                             sheetUpload["CONSEQUENCE_DEFECT"] = sheetUpload["CONSEQUENCE_DEFCT"]
#                             sheetUpload.drop(['CONSEQUENCE_DEFCT'], axis=1, inplace = True)
#                             sheetUpload.to_sql(formDict.get(sheet),con=engine, if_exists='append', index=False)
#                         else:
#                             sheetUpload.to_sql(formDict.get(sheet),con=engine, if_exists='append', index=False)
                    
#                         # df_update.to_sql(table_name, con=engine, if_exists='append', index=False)
#                         # print('upload SUCCESS!: ', sheet)
#                         logger_bp_admin.info(f"upload SUCCESS!:: {sheet}")
#                         sheet_upload_status.append(f"{sheet}: success")
#                     except IndexError:
#                         logger_bp_admin.info(f"except IndexError:: {IndexError}")
#                         # return redirect(url_for('bp_admin.database_page',legend=legend,
#                         #     tableNamesList=tableNamesList, sheetNames=sheetNames))
#                         sheet_upload_status.append(f"{sheet}: fail")
#                     except:
#                         logger_bp_admin.info(f"except another error:: {sheet}")
#                         # os.remove(os.path.join(current_app.config['DIR_DB_FILES_TEMPORARY'], uploadFileName))
#                         sheet_upload_status.append(f"{sheet}: fail")

#                         # flash(f"""Problem uploading {sheet} table. Check for 1)uniquness with id or RECORD_ID 2)date columns
#                         #     are in a date format in excel.""", 'warning')
#                         # return redirect(url_for('bp_admin.database_page',legend=legend,
#                         #     tableNamesList=tableNamesList, sheetNames=sheetNames))
#                     #clear files_temp folder
#                 for file in os.listdir(current_app.config['DIR_DB_FILES_TEMPORARY']):
#                     os.remove(os.path.join(current_app.config['DIR_DB_FILES_TEMPORARY'], file))
#                 status_message =""
#                 for sheet_message in sheet_upload_status:
#                     status_message = status_message +sheet_message + ",\n"
#                 flash(f'Table sheet status: {status_message}', 'info')
#                 return redirect(url_for('bp_admin.database_page',legend=legend,
#                     tableNamesList=tableNamesList, sheetNames=sheetNames))

                
#             elif file_type=='text':
#                 zipfile.ZipFile(uploaded_file).extractall(path=current_app.config['DIR_DB_FILES_TEMPORARY'])
                
                
#                 text_file_name=[x for x in os.listdir(current_app.config['DIR_DB_FILES_TEMPORARY']) if x[-4:]=='.txt'][0]
#                 limit_upload_flag=formDict.get('limit_upload_flag')
                
#                 flash_message=load_database_util(text_file_name, limit_upload_flag)
                
                
                
#                 for file in os.listdir(current_app.config['DIR_DB_FILES_TEMPORARY']):
#                     os.remove(os.path.join(current_app.config['DIR_DB_FILES_TEMPORARY'], file))

                    
#                 flash(flash_message[0], flash_message[1])

#                 return redirect(url_for('bp_admin.database_page',legend=legend))

    
#     if file_type=='excel':
#         return render_template('admin/database_upload.html',legend=legend,tableNamesList=tableNamesList,
#                     sheetNames=sheetNames, uploadFileName=uploadFileName,
#                     # uploadFlag=uploadFlag,
#                     file_type=file_type)
#     else:
#         return render_template('admin/database_upload.html',legend=legend,
#                     uploadFileName=uploadFileName,
#                     # uploadFlag=uploadFlag,
#                     file_type=file_type,limit_upload_flag=limit_upload_flag)










################
### VERY OLD ###
################



# @bp_admin.route('/database_delete_data', methods=["GET","POST"])
# @login_required
# def database_delete_data():

#     return redirect(request.referrer)



# @bp_admin.route('/admin_db_download', methods = ['GET', 'POST'])
# @login_required
# def admin_db_download():
#     logger_bp_admin.info('- in bp_admin_db_download -')
#     logger_bp_admin.info(f"current_user.admin: {current_user.admin}")

#     if not current_user.admin:
#         return redirect(url_for('bp_bp_main.home'))

#     metadata = Base.metadata
#     db_table_list = [table for table in metadata.tables.keys()]

#     csv_dir_path = os.path.join(current_app.config.get('PROJECT_RESOURCES'), 'db_backup')

#     if request.method == "POST":
#         formDict = request.form.to_dict()
#         # print(f"- search_rincons POST -")
#         # print("formDict: ", formDict)

#         # craete folder to save
#         if not os.path.exists(os.path.join(os.environ.get('PROJECT_RESOURCES'),"db_backup")):
#             os.makedirs(os.path.join(os.environ.get('PROJECT_RESOURCES'),"db_backup"))


#         db_table_list = []
#         for key, value in formDict.items():
#             if value == "db_table":
#                 db_table_list.append(key)
      
#         db_tables_dict = {}
#         for table_name in db_table_list:
#             base_query = sess.query(metadata.tables[table_name])
#             df = pd.read_sql(text(str(base_query)), engine.connect())

#             # fix table names
#             cols = list(df.columns)
#             for col in cols:
#                 if col[:len(table_name)] == table_name:
#                     df = df.rename(columns=({col: col[len(table_name)+1:]}))

#             # Users table convert password from bytes to strings
#             if table_name == 'users':
#                 df['password'] = df['password'].str.decode("utf-8")


#             db_tables_dict[table_name] = df
#             db_tables_dict[table_name].to_csv(os.path.join(csv_dir_path, f"{table_name}.csv"), index=False)
        
#         shutil.make_archive(csv_dir_path, 'zip', csv_dir_path)

#         return redirect(url_for('bp_admin.download_db_tables_as_csv'))
    
#     return render_template('bp_admin/admin_db_download.html', db_table_list=db_table_list )

# @bp_admin.route("/download_db_tables_as_csv", methods=["GET","POST"])
# @login_required
# def download_db_tables_as_csv():
#     return send_from_directory(os.path.join(current_app.config['PROJECT_RESOURCES']),'db_backup.zip', as_attachment=True)



# @bp_admin.route('/admin_db_upload_single_file', methods = ['GET', 'POST'])
# @login_required
# def admin_db_upload_single_file():
#     logger_bp_admin.info('- in bp_admin_db_upload_single_file -')
#     logger_bp_admin.info(f"current_user.admin: {current_user.admin}")

#     if not current_user.admin:
#         return redirect(url_for('bp_bp_main.home'))

#     metadata = Base.metadata
#     db_table_list = [table for table in metadata.tables.keys()]
#     csv_dir_path_upload = os.path.join(current_app.config.get('PROJECT_RESOURCES'), 'db_upload')

#     if request.method == "POST":
#         formDict = request.form.to_dict()
#         # print(f"- search_rincons POST -")
#         # print("formDict: ", formDict)

#         requestFiles = request.files

#         # print("requestFiles: ", requestFiles)

#         # craete folder to store upload files
#         if not os.path.exists(os.path.join(os.environ.get('PROJECT_RESOURCES'),"db_upload")):
#             os.makedirs(os.path.join(os.environ.get('PROJECT_RESOURCES'),"db_upload"))
        

#         csv_file_for_table = request.files.get('csv_table_upload')
#         csv_file_for_table_filename = csv_file_for_table.filename

#         logger_bp_admin.info(f"-- Get CSV file name --")
#         logger_bp_admin.info(f"--  {csv_file_for_table_filename} --")

#         ## save to static rincon directory
#         path_to_uploaded_csv = os.path.join(csv_dir_path_upload,csv_file_for_table_filename)
#         csv_file_for_table.save(path_to_uploaded_csv)

#         print(f"-- table to go to: { formDict.get('existing_db_table_to_update')}")

#         return redirect(url_for('bp_admin.upload_table', table_name = formDict.get('existing_db_table_to_update'),
#             path_to_uploaded_csv=path_to_uploaded_csv))


#     return render_template('admin/admin_db_upload_single_file.html', db_table_list=db_table_list)


# @bp_admin.route('/upload_table/<table_name>', methods = ['GET', 'POST'])
# @login_required
# def upload_table(table_name):
#     logger_bp_admin.info('- in upload_table -')
#     logger_bp_admin.info(f"current_user.admin: {current_user.admin}")
#     path_to_uploaded_csv = request.args.get('path_to_uploaded_csv')

#     if not current_user.admin:
#         return redirect(url_for('bp_bp_main.home'))

#     # Get Table Column names from the database corresponding to the running webiste/app
#     metadata = Base.metadata
#     existing_table_column_names = metadata.tables[table_name].columns.keys()

#     # Get column names from the uploaded csv
#     df = pd.read_csv(path_to_uploaded_csv)

#     if 'time_stamp_utc' in df.columns:
#         try:
#             df['time_stamp_utc'] = pd.to_datetime(df['time_stamp_utc'], format='%d/%m/%Y %H:%M')
#         # except ValueError:
#         #     df['time_stamp_utc'] = pd.to_datetime(df['time_stamp_utc'], format='%d/%m/%Y %H:%M:%S')
#         except:
#             df = pd.read_csv(path_to_uploaded_csv, parse_dates=['time_stamp_utc'])


    

#     replacement_data_col_names = list(df.columns)

#     # Match column names between the two tables
#     match_cols_dict = {}
#     for existing_db_column in existing_table_column_names:
#         try:
#             index = replacement_data_col_names.index(existing_db_column)
#             match_cols_dict[existing_db_column] = replacement_data_col_names[index]
#         except ValueError:
#             match_cols_dict[existing_db_column] = None


#     if request.method == "POST":
#         formDict = request.form.to_dict()
#         # print(f"- search_rincons POST -")
#         # print("formDict: ", formDict)

#         # NOTE: upload data to existing database
#         ### formDict (key) is existing databaes column name
#         # existing_names_list = [existing for existing, update in formDict.items() if update != 'true' ]
        
#         # check for default values and remove from formDict
#         set_default_value_dict = {}
#         for key, value in formDict.items():
#             if key[:len("default_checkbox_")] == "default_checkbox_":
#                 set_default_value_dict[value] = formDict.get(value)
        

#         # Delete elements from dictionary
#         for key, value in set_default_value_dict.items():
#             del formDict[key]
#             checkbox_key = "default_checkbox_" + key
#             del formDict[checkbox_key]




#         print("- formDict adjusted -")
#         print(formDict)

#         existing_names_list = []
#         for key, value in formDict.items():
#             if value != 'true':
#                 existing_names_list.append(key)
            
            
#         df_update = pd.DataFrame(columns=existing_names_list)

#         # value is the new data (aka the uploaded csv file column)
#         for exisiting, replacement in formDict.items():
#             if not replacement in ['true','']:
#                 # print(replacement)
#                 df_update[exisiting]=df[replacement].values

#         # Add in columns with default values
#         for column_name, default_value in set_default_value_dict.items():
#             if column_name == 'time_stamp_utc': 
#                 df_update[column_name] = datetime.utcnow()
#             else:
#                 df_update[column_name] = default_value

        
#         # remove existing users from upload
#         # NOTE: There needs to be a user to upload data
#         if table_name == 'users':
#             print("--- Found users table ---")
#             existing_users = sess.query(Users).all()
#             list_of_emails_in_db = [i.email for i in existing_users]
#             for email in list_of_emails_in_db:
#                 df_update.drop(df_update[df_update.email== email].index, inplace = True)
#                 print(f"-- removeing {email} from upload dataset --")
        

#             for index in range(1,len(df_update)+1):
#                 df_update.loc[index, 'password'] = df_update.loc[index, 'password'].encode()
#                 # print(" ****************** ")
#                 # print(f"- encoded row for {df_update.loc[index, 'email']} -")
#                 # print(" ****************** ")


#         df_update.to_sql(table_name, con=engine, if_exists='append', index=False)

#         flash(f"{table_name} update: successful!", "success")

#         # print("request.path: ", request.path)
#         # print("request.full_path: ", request.full_path)
#         # print("request.script_root: ", request.script_root)
#         # print("request.base_url: ", request.base_url)
#         # print("request.url: ", request.url)
#         # print("request.url_root: ", request.url_root)
#         # print("______")


#         # return redirect(request.url)
#         return redirect(url_for('bp_admin.admin_db_upload_single_file'))


    
#     return render_template('admin/upload_table.html', table_name=table_name, 
#         match_cols_dict = match_cols_dict,
#         existing_table_column_names=existing_table_column_names,
#         replacement_data_col_names = replacement_data_col_names)





